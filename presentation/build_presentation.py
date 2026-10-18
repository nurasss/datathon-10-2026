"""Build an editable PowerPoint and matching PDF. No bank footage is embedded."""
from pathlib import Path
import json
import math

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path(__file__).resolve().parent
ROOT = OUT.parent
METRICS = ROOT / 'fv-core/artifacts/task8/normal/metrics.json'
data = json.loads(METRICS.read_text())
s = data['summary']
assert (s['tp'], s['fp'], s['fn'], s['tn']) == (11, 0, 10, 19), 'Update slide narrative for new metrics'

BG = '101820'
PANEL = '1B2834'
WHITE = 'F4F3EC'
MUTED = 'A8B8C6'
BLUE = '6BAAFF'
LIME = 'D2F280'
LINE = '354654'
AMBER = 'F2BE75'
W, H = 13.333333, 7.5

fonts = Path('/System/Library/Fonts/Supplemental')
pdfmetrics.registerFont(TTFont('Arial', str(fonts / 'Arial.ttf')))
pdfmetrics.registerFont(TTFont('Arial-Bold', str(fonts / 'Arial Bold.ttf')))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'ForteVision — локальный AI-контроль кассовой дисциплины'
prs.core_properties.subject = 'ForteBank · Datathon 26 · питч на 5 минут'
prs.core_properties.author = 'ForteVision'
pdf = canvas.Canvas(str(OUT / 'ForteVision_pitch.pdf'), pagesize=(W * 72, H * 72))
pdf.setTitle(prs.core_properties.title)
notes = []
slide = None


def rgb(c):
    return RGBColor.from_string(c)


def rect(x, y, w, h, color, border=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    if border:
        sh.line.color.rgb = rgb(border)
        sh.line.width = Pt(.7)
    else:
        sh.line.fill.background()
    pdf.setFillColor('#' + color)
    pdf.setStrokeColor('#' + (border or color))
    pdf.setLineWidth(.7 if border else 0)
    pdf.rect(x * 72, (H - y - h) * 72, w * 72, h * 72, fill=1, stroke=int(bool(border)))


def wrap(text, width, size, bold):
    font = 'Arial-Bold' if bold else 'Arial'
    lines = []
    for para in text.split('\n'):
        line = ''
        for word in para.split():
            candidate = (line + ' ' + word).strip()
            if pdfmetrics.stringWidth(candidate, font, size) > width * 72 - 5 and line:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


def text(x, y, w, content, size=20, color=WHITE, bold=False, leading=1.2):
    lines = wrap(content, w, size, bold)
    h = len(lines) * size * leading / 72 + .06
    assert y + h < 7.46, (content, y, h)
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = 'Arial', Pt(size), bold
        p.font.color.rgb = rgb(color)
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size * leading)
    pdf.setFillColor('#' + color)
    pdf.setFont('Arial-Bold' if bold else 'Arial', size)
    for i, line in enumerate(lines):
        pdf.drawString(x * 72, (H - y) * 72 - size * .92 - i * size * leading, line)
    return h


def new(kicker, title, note, seconds=0):
    global slide
    if slide is not None:
        pdf.showPage()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    rect(0, 0, W, H, BG)
    rect(.6, .47, .1, .14, LIME)
    text(.82, .43, 11.8, kicker.upper(), 11, MUTED, True)
    if title:
        text(.65, .99, 12.05, title, 32, WHITE, True)
    rect(.65, 6.99, 12.03, .012, LINE)
    text(.65, 7.12, 10.9, 'FORTEVISION  /  FORTEBANK DATATHON 26  /  ОКТЯБРЬ 2026', 8.5, MUTED)
    text(12.12, 7.1, .6, f'{len(prs.slides):02d}', 10, MUTED)
    slide.notes_slide.notes_text_frame.text = note
    notes.append((kicker, note, seconds))


def panel(x, y, w, h, label, title, body, accent=BLUE):
    rect(x, y, w, h, PANEL)
    rect(x, y, .045, h, accent)
    text(x + .24, y + .24, w - .48, label, 11, accent, True)
    text(x + .24, y + .65, w - .48, title, 22, WHITE, True)
    text(x + .24, y + 1.35, w - .48, body, 17, MUTED)


new('01 / Локальный AI-сервис', '',
    'ForteVision — локальный сервис для проверки видеозаписей кассы. Мы не заменяем контролёра: '
    'находим визуальные признаки возможного нарушения и связываем их с моментом на видео, '
    'кадрами и нормативным источником. Представляем работающий MVP, а не только концепцию.', 20)
text(.65, 1.27, 12, 'ForteVision', 70, WHITE, True)
text(.69, 2.57, 10.8, 'От видеозаписи кассы\nк проверяемому событию.', 35, WHITE, True)
text(.72, 4.27, 9.1, 'Таймкод. Кадр-доказательство. Норма РК.\nОкончательное решение — за человеком.', 23, MUTED)
rect(.72, 5.66, 11.9, .7, PANEL)
text(.96, 5.85, 11.45, 'ЛОКАЛЬНОЕ ВЫПОЛНЕНИЕ     /     ВЕБ-ИНТЕРФЕЙС     /     API ДЛЯ ЖЮРИ', 14, LIME, True)

new('02 / Задача банка', 'Контролёру нужен фрагмент, а не весь ролик.',
    'При ручной проверке контролёру нужно просмотреть запись, найти событие, зафиксировать время '
    'и понять, какое правило применимо. Наш продукт собирает эту цепочку в одном интерфейсе. '
    'Бизнес-гипотеза — сократить поиск и подготовку материалов для проверки. Экономию времени '
    'пока не измеряли и не заявляем числом.', 25)
panel(.65, 2.1, 3.87, 3.1, '01 / НАЙТИ', 'Что произошло?', 'Выделить визуально наблюдаемое действие среди обычной работы кассира.')
panel(4.73, 2.1, 3.87, 3.1, '02 / ПРОВЕРИТЬ', 'Где доказательство?', 'Открыть нужный момент и сравнить вывод с кадрами исходного видео.')
panel(8.81, 2.1, 3.87, 3.1, '03 / ОБОСНОВАТЬ', 'Какое правило?', 'Увидеть источник, пункт и ограничения интерпретации.', LIME)
text(.7, 5.71, 11.9, 'Ценность: меньше поиска — больше времени на проверку контекста.', 23, WHITE, True)

new('03 / Пользовательский сценарий', 'Загрузить → дождаться → проверить.',
    'Пользователь загружает MOV через веб-интерфейс. Сервер создаёт задачу и показывает статус '
    'обработки. На выходе — один из двух вердиктов. Для каждого найденного события доступны тип, '
    'таймкод, описание и до трёх JPEG-кадров. Обоснование содержит нормативную ссылку. '
    'Схема на слайде иллюстрирует состав результата, это не скриншот отдельного детектирования.', 30)
for i, (a, b) in enumerate([
    ('Загрузка .mov', 'Через браузер или API'),
    ('processing', 'Статус фоновой задачи'),
    ('done / failed', 'Результат или явная ошибка'),
]):
    x = .65 + i * 4.08
    rect(x, 2.0, 3.87, 1.12, PANEL)
    text(x + .2, 2.19, 3.48, a, 22, WHITE, True)
    text(x + .2, 2.67, 3.48, b, 13, MUTED)
rect(.65, 3.48, 12.03, 2.95, PANEL, LINE)
text(.92, 3.74, 5.5, 'Карточка возможного нарушения', 22, BLUE, True)
text(.92, 4.35, 5.5, 'Тип + таймкод\nНаблюдаемое действие\nПравило, источник и оговорки', 20, WHITE, leading=1.45)
for i, label in enumerate(['Кадр до', 'Момент', 'Кадр после']):
    x = 7.0 + i * 1.75
    rect(x, 4.25, 1.55, 1.2, BG, LINE)
    rect(x + .2, 4.55, 1.15, .035, BLUE if i == 1 else LINE)
    rect(x + .2, 4.83, .7, .035, LINE)
    text(x, 5.65, 1.65, label, 12, MUTED)
text(.93, 6.12, 6, 'Схема результата • не фактический кейс', 11, MUTED)

new('04 / Архитектура', 'Весь анализ — внутри локального сервиса.',
    'FastAPI принимает файл и сохраняет задачу в SQLite. Отдельные процессы выполняют '
    'анализ на CPU: NanoDet в ONNX Runtime, трекинг людей, анализ визуальных состояний и '
    'типизированные правила. Редкий проход выделяет кандидатные окна, плотный уточняет '
    'события. Evidence извлекается из исходного видео. Облачные AI API не используются; '
    'лица не идентифицируем и звук не анализируем. Публичный туннель даёт доступ к сервису, '
    'но не выполняет AI-инференс.', 40)
steps = [('01', 'Приём видео', 'FastAPI + SQLite'), ('02', 'Зрение', 'NanoDet + ONNX CPU'),
         ('03', 'События', 'Треки + состояния'), ('04', 'Правила', 'Проверенный реестр')]
for i, (n, title, sub) in enumerate(steps):
    x = .65 + i * 3.06
    rect(x, 2.24, 2.85, 2.15, PANEL)
    text(x + .2, 2.44, 2.45, n, 16, LIME, True)
    text(x + .2, 3.03, 2.45, title, 22, WHITE, True)
    text(x + .2, 3.67, 2.45, sub, 13, MUTED)
text(.7, 4.88, 11.9, 'Редкий проход → кандидатные окна → плотная проверка → JSON + кадры', 21, BLUE, True)
text(.7, 5.66, 11.8, 'Без облачных AI API. Без анализа звука. Без идентификации лиц.\nПосле установки зависимостей и весов интернет для анализа не нужен.', 18, MUTED)

new('05 / Обнаруживаемые события', 'Правило в коде ≠ доказанное качество.',
    'В коде предусмотрены восемь основных типов и дополнительный уход во время обслуживания. '
    'Но мы разделяем наличие правила и измеренную работоспособность. Последний полный прогон '
    'подтверждает находки четырёх типов: ценности, экран, посторонний и дверь. Сейф, личные вещи '
    'и ручной пересчёт в этом прогоне пропущены полностью. Для телефона нет положительных '
    'примеров, поэтому оценивать его качество нельзя.', 30)
text(.7, 2.05, 5.7, 'Есть находки в полном прогоне', 20, LIME, True)
text(7.0, 2.05, 5.55, 'Нужна доработка / оценка', 20, AMBER, True)
left = ['Ценности оставлены при уходе', 'Уход без блокировки экрана', 'Вход человека в служебную зону', 'Дверь оставлена открытой']
right = ['Открытый сейф без присмотра', 'Личные вещи / еда / напитки', 'Ручной пересчёт наличных', 'Использование телефона*']
for i, (a, b) in enumerate(zip(left, right)):
    y = 2.75 + i * .7
    rect(.7, y + .1, .09, .09, LIME)
    text(.97, y, 5.6, a, 18)
    rect(7.0, y + .1, .09, .09, AMBER)
    text(7.27, y, 5.35, b, 18, MUTED)
text(.7, 6.17, 11.9, '* Для телефона нет положительных примеров. Наличие находок не означает надёжность на новом ракурсе.', 12, MUTED)

new('06 / Обоснованность', 'Ссылка на норму — из реестра, не из генерации.',
    'Нормативная база включает акты Республики Казахстан: постановление НБРК 231 о кассовых '
    'операциях, НБРК 14 об устройстве и охране помещений и АРРФР 53 об информационной '
    'безопасности. Система берёт цитату и ссылку из фиксированного реестра. Важен и предел '
    'вывода: например, по видео нельзя установить полномочия вошедшего. Поэтому результат — '
    'признак для проверки, не юридический вердикт. Экран особенно спорен: норма об автоблокировке '
    'через 30 минут не доказывает нарушение при кратком уходе.', 35)
panel(.65, 2.1, 3.87, 2.8, 'НАБЛЮДЕНИЕ', 'Видимый факт', 'Человек пересёк вход и прошёл в служебную зону.')
panel(4.73, 2.1, 3.87, 2.8, 'ГИПОТЕЗА', 'Возможный риск', 'Может указывать на нарушение режима изоляции кассы.')
panel(8.81, 2.1, 3.87, 2.8, 'ОГРАНИЧЕНИЕ', 'Не доказано', 'Наличие или отсутствие служебных полномочий вошедшего.', LIME)
text(.7, 5.3, 11.9, 'НБРК №231  /  НБРК №14  /  АРРФР №53', 24, WHITE, True)
text(.7, 5.92, 11.9, 'Документ + пункт + цитата + ссылка на adilet.zan.kz. Решение принимает проверяющий.', 16, MUTED)

new('07 / Реальные результаты', 'Находим 11 из 21 видео с нарушениями.',
    'На 40 доступных размеченных видео система правильно определила 11 положительных и '
    'все 19 отрицательных. Десять положительных видео пропущены. Это 75 процентов правильных '
    'бинарных вердиктов, но полнота обнаружения только 52,4 процента. Ноль ложных тревог '
    'относится к этим 19 нормальным роликам, не ко всем будущим данным. Медианная ошибка '
    'таймкода 1,4 секунды, максимальная — 17. Это внутренний прогон на известном датасете, '
    'не независимая оценка жюри.', 45)
for x, value, label, sub, col in [
    (.7, '52,4%', 'Полнота обнаружения', '11 из 21 положительного видео', BLUE),
    (4.87, '0 / 19', 'Ложные тревоги', 'На нормальных видео прогона', LIME),
    (9.04, '1,4 с', 'Ошибка таймкода', 'Медиана; максимум — 17 с', WHITE),
]:
    text(x, 2.16, 3.68, value, 48, col, True)
    text(x, 3.2, 3.68, label, 19, WHITE, True)
    text(x, 3.78, 3.68, sub, 14, MUTED)
text(.7, 4.62, 11.9, '40 видео: 21 с нарушением + 19 без нарушения', 20, WHITE, True)
for i, (val, col) in enumerate([(11, BLUE), (10, AMBER), (19, LIME)]):
    prev = [0, 11, 21][i]
    rect(.7 + prev / 40 * 11.9, 5.2, val / 40 * 11.9 - .025, .42, col)
text(.7, 5.83, 11.9, '11 найдены     /     10 пропущены     /     19 верных отрицательных', 17, MUTED)
text(.7, 6.39, 11.9, 'Известный датасет, не независимый тест. Источник: artifacts/task8/normal/metrics.json.', 11, MUTED)

new('08 / Готовность и развитие', 'Рабочий MVP. Следующий приоритет — пропуски.',
    'Продуктовый контур уже работает: интерфейс, API с ключом, очередь, таймауты и кадры. '
    'В последнем локальном прогоне анализ занимал от 4,45 до 36,94 секунды на видео, '
    'но это не включает сетевую загрузку и очередь. Для сдачи нужен сквозной публичный '
    'замер в пределах пяти минут. Развитие: сначала работа с пропущенными типами, '
    'затем оценка на невиденных ракурсах и обновление документации. Это планы, '
    'а не уже достигнутые показатели.', 40)
panel(.65, 2.07, 5.85, 3.9, 'УЖЕ ЕСТЬ', 'Полный сервисный контур',
      'Веб-интерфейс + API с ключом\nФоновая обработка и watchdog\nРеальные кадры из исходника\nЛокальный CPU-инференс', LIME)
panel(6.75, 2.07, 5.93, 3.9, 'СЛЕДУЮЩИЙ ЭТАП', 'Улучшить зрение и проверку',
      '1. Сейф, предметы, пересчёт\n2. Новые ракурсы без подстройки\n3. Публичный тест до 5 минут\n4. Актуальные документы сдачи', BLUE)
text(.7, 6.28, 11.9, 'Локальный анализ: 4,45–36,94 с/видео. Без сетевой загрузки и ожидания очереди.', 13, MUTED)

new('09 / Демонстрация', 'Проверяемый сигнал вместо чёрного ящика.',
    'В завершение предлагаем открыть сервис и загрузить MOV. Показать статус задачи, '
    'затем вердикт, таймкод, кадры и источник. Ключи доступа передаются жюри отдельно, '
    'в презентации их нет. Если сервис или туннель недоступны, не выдаём схему за живое '
    'демо. Главное сообщение: инженерный контур готов, качество измерено честно, '
    'следующий фокус — полнота и переносимость на другие ракурсы.', 35)
text(.7, 2.17, 10.9, 'Видео → событие → доказательство → норма', 32, WHITE, True)
text(.7, 3.19, 10.9, 'Покажем загрузку, статус обработки\nи карточку найденного события.', 25, MUTED)
rect(.7, 4.59, 11.9, 1.4, PANEL)
text(.97, 4.81, 11.3, 'ОТКРЫТЬ СЕРВИС', 12, LIME, True)
text(.97, 5.24, 11.3, 'https://craig-taxaceous-marylouise.ngrok-free.dev', 19, BLUE)
text(.7, 6.32, 11.9, 'Ключ доступа — отдельно. Доступность URL зависит от локального сервера и туннеля.', 12, MUTED)

new('Приложение A / Для вопросов жюри', 'Точность по типам: где система уже полезна.',
    'Резервный слайд. Метрики по типам считаются отдельно от бинарного вердикта. '
    'Две ошибочные классификации как ценности возникли на положительных видео с другими '
    'типами и не стали бинарными ложными тревогами. Событий 22, положительных видео 21: '
    'одно видео содержит два события. Телефон оценить нельзя: нет положительных примеров.')
headers = [('Тип события', .85, 6.5), ('Найдено / всего', 7.6, 2.4), ('Ложные типы', 10.5, 2.0)]
for label, x, width in headers:
    text(x, 1.98, width, label, 15, BLUE, True)
rows = [('Ценности при уходе', '2 / 3', '2'), ('Незаблокированный экран', '2 / 2', '0'),
        ('Открытый сейф', '0 / 2', '0'), ('Личные вещи / еда / напитки', '0 / 3', '0'),
        ('Посторонний в кассе', '4 / 6', '0'), ('Открытая дверь', '2 / 2', '0'),
        ('Ручной пересчёт', '0 / 4', '0'), ('Телефон', 'нет примеров', '0')]
for i, row in enumerate(rows):
    y = 2.43 + i * .43
    if i % 2 == 0:
        rect(.7, y - .03, 11.9, .42, PANEL)
    for (label, x, width), val in zip(headers, row):
        text(x, y, width, val, 15, WHITE)
text(.85, 6.29, 11.5, '22 размеченных события в 21 положительном видео. Ложный тип ≠ ложный бинарный вердикт.', 12, MUTED)

new('Приложение B / Интеграция', 'Контракт API соответствует кейсу.',
    'Резервный технический слайд. POST принимает multipart file и возвращает job_id. '
    'GET возвращает processing, done или failed. В готовом ответе только два вердикта, '
    'список событий и JPEG в base64. Название нормы и ссылка включены в description; '
    'расширенный details endpoint отдаёт структуру правила отдельно. Задача завершается '
    'по таймауту 280 секунд после приёма файла, включая очередь. Это защита от зависания, '
    'а не гарантия успешного анализа. Сохранённый тестовый отчёт — 139 пройденных тестов.')
text(.72, 2.08, 11.9, 'POST /api/v1/analyze', 24, BLUE, True)
text(.72, 2.64, 11.9, 'X-API-Key + multipart file (.mov) → job_id', 19, MUTED)
text(.72, 3.26, 11.9, 'GET /api/v1/jobs/{job_id}', 24, BLUE, True)
text(.72, 3.83, 11.9, 'processing → done / failed', 19, MUTED)
rect(.72, 4.53, 11.88, 1.16, PANEL)
text(.97, 4.73, 11.3, 'verdict • violations[] • type • timecode_sec\ndescription • evidence_frames[]: timestamp_sec + image_base64', 18, WHITE)
text(.72, 6.04, 11.9, 'До 100 MiB • watchdog 280 с после приёма • неверный ключ: 401\n139 passed в сохранённом отчёте тестов; это не метрика качества распознавания.', 14, MUTED)

prs.save(OUT / 'ForteVision_pitch.pptx')
pdf.save()
md = ['# ForteVision — текст выступления', '',
      '9 основных слайдов + 2 резервных. Плановое время: 5 минут. '
      'Данные: последний найденный полный прогон task8/normal. '
      'В презентацию не включены видео банка, кадры датасета или ключи доступа.', '']
for i, (title, note, sec) in enumerate(notes, 1):
    md += [f'## {i}. {title}' + (f' — {sec} с' if sec else ' — резерв'), '', note, '']
md += ['## Источники и оговорки', '',
       '- `fv-core/artifacts/task8/normal/metrics.json` — бинарные и типовые метрики, время анализа.',
       '- `fv-core/artifacts/task8/pytest_selected.txt` — сохранённый результат 139 passed.',
       '- `forte-vision/apps/api/app/main.py` — API и UI.',
       '- `forte-vision/packages/rules/rules/registry.json` — реестр нормативных ссылок.',
       '- `forte-vision/configs/inference.default.json` — текущий конфиг.',
       '- НБРК №231: https://adilet.zan.kz/rus/docs/V1900019680',
       '- НБРК №14: https://adilet.zan.kz/rus/docs/V2000020075',
       '- АРРФР №53: https://adilet.zan.kz/rus/docs/V2600038336',
       '- Цитаты заимствованы из реестра проекта; самостоятельная правовая экспертиза не проводилась.',
       '- 75% accuracy = (11 + 19) / 40; recall = 11 / 21; precision = 11 / 11. '
       '100% precision в этом прогоне не означает безошибочность на новых видео.',
       '- Таймкод оценивается на бинарно положительных совпадениях; это не отдельная '
       'гарантия правильного типа или полноты событий.',
       '- Никакие целевые будущие метрики не заявляются как достигнутые.',
       '- Перед выступлением проверить публичный URL и иметь отдельный ключ доступа.', '']
(OUT / 'SPEAKER_NOTES.md').write_text('\n'.join(md))
(OUT / 'README.md').write_text(
    '# ForteVision — презентация\n\n'
    '- `ForteVision_pitch.pptx` — редактируемый PowerPoint 16:9 с заметками докладчика.\n'
    '- `ForteVision_pitch.pdf` — PDF для просмотра и отправки.\n'
    '- `SPEAKER_NOTES.md` — текст 5-минутного выступления и источники.\n'
    '- Слайды 1–9 — питч (300 секунд); 10–11 — ответы на вопросы.\n'
    '- Банковские кадры, исходные видео и секретные ключи не включены.\n\n'
    'Пересборка на этом Mac:\n\n```bash\n'
    '/private/var/folders/ws/2yz96w190q34htgtvflq4bm00000gn/T/opencode/forte-slides-env/bin/python '
    'presentation/build_presentation.py\n```\n'
)
print(f'Created {len(prs.slides)} slides; main pitch: {sum(n[2] for n in notes)} seconds')
print(OUT / 'ForteVision_pitch.pptx')
print(OUT / 'ForteVision_pitch.pdf')
